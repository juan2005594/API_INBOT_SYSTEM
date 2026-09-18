from io import BytesIO

from django.conf import settings
from django.core.mail import EmailMessage
from django.utils import timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..models import Venta


def _docx_to_bytes(doc: Document) -> bytes:
    """
    Convierte un docx generado con python-docx a bytes para adjuntarlo en un correo.
    """
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()


def _build_factura_docx(venta: Venta, detalles, nombre_cliente=None, documento_cliente=None) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('FACTURA DE VENTA')
    run.bold = True

    doc.add_paragraph(f'ID Venta: {venta.id}')
    doc.add_paragraph(f'Fecha: {venta.fecha_venta.strftime("%Y-%m-%d %H:%M:%S")}')
    doc.add_paragraph(f'Vendedor: {venta.vendedor}')
    doc.add_paragraph(f'Método de pago: {venta.metodo_pago}')
    doc.add_paragraph(f'Total: {venta.total:.2f}')

    doc.add_paragraph()
    doc.add_paragraph('Datos del cliente')
    doc.add_paragraph(f'Nombre: {nombre_cliente or "N/D"}')
    doc.add_paragraph(f'Documento: {documento_cliente or "N/D"}')

    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = 'Producto'
    hdr[1].text = 'Cantidad'
    hdr[2].text = 'Precio Unit.'
    hdr[3].text = 'Subtotal'

    for d in detalles:
        row = table.add_row().cells
        row[0].text = d.producto.nombre
        row[1].text = str(d.cantidad)
        row[2].text = f'{d.precio_unitario:.2f}'
        row[3].text = f'{d.subtotal:.2f}'

    return _docx_to_bytes(doc)


def _build_comprobante_docx(venta: Venta, detalles) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('COMPROBANTE DE PAGO')
    run.bold = True

    doc.add_paragraph()
    doc.add_paragraph('Este comprobante se genera automáticamente al registrar la venta.')
    doc.add_paragraph(f'ID Venta: {venta.id}')
    doc.add_paragraph(f'Fecha y hora de emisión: {timezone.localtime().strftime("%Y-%m-%d %H:%M:%S")}')
    doc.add_paragraph(f'Vendedor: {venta.vendedor}')
    doc.add_paragraph(f'Método de pago: {venta.metodo_pago}')
    doc.add_paragraph(f'Total pagado: {venta.total:.2f}')

    doc.add_paragraph()
    doc.add_paragraph('Resumen de productos')
    for d in detalles:
        doc.add_paragraph(f'- {d.producto.nombre}: {d.cantidad} unidades')

    return _docx_to_bytes(doc)


def send_venta_invoice_email(venta_id: int, correo_cliente: str, nombre_cliente=None, documento_cliente=None) -> None:
    """
    Envía por correo un DOCX con factura y otro DOCX con comprobante.
    Funciona como "mejor esfuerzo": si falta la configuración SMTP, no falla el flujo de registro.
    """
    # Verificación de configuración SMTP.
    if not getattr(settings, 'EMAIL_HOST_USER', '') or not getattr(settings, 'EMAIL_HOST_PASSWORD', ''):
        print('DEBUG: No se configuró EMAIL_HOST_USER/EMAIL_HOST_PASSWORD. No se enviará el correo.')
        return

    venta = (
        Venta.objects
        .prefetch_related('detalles__producto')
        .get(id=venta_id)
    )
    detalles = list(venta.detalles.select_related('producto').all())

    factura_bytes = _build_factura_docx(venta, detalles, nombre_cliente=nombre_cliente, documento_cliente=documento_cliente)
    comprobante_bytes = _build_comprobante_docx(venta, detalles)

    subject = f'Factura de venta - ID {venta.id}'
    body = (
        f'Hola {nombre_cliente or "cliente"},\n\n'
        f'Adjuntamos la factura y el comprobante de pago de tu venta (ID: {venta.id}).\n\n'
        'Gracias por tu compra.'
    )

    email = EmailMessage(
        subject=subject,
        body=body,
        from_email=getattr(settings, 'DEFAULT_FROM_EMAIL', settings.EMAIL_HOST_USER),
        to=[correo_cliente],
    )

    email.attach(
        filename=f'factura_venta_{venta.id}.docx',
        content=factura_bytes,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    email.attach(
        filename=f'comprobante_pago_{venta.id}.docx',
        content=comprobante_bytes,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )

    try:
        email.send(fail_silently=False)
        print(f'DEBUG: Correo de factura enviado a {correo_cliente} para venta {venta_id}')
    except Exception as e:
        # Mejor esfuerzo: el flujo de registro ya pasó; solo registramos el error.
        print(f'DEBUG: Error enviando correo para venta {venta_id}: {e}')
        return

